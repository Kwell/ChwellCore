"""Excel import and compiled C# contracts; also run against the installed tool."""
import argparse
import copy
import datetime
import importlib.util
import json
import pathlib
import shutil
import subprocess
import sys
import tempfile
import unittest

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--tool', type=pathlib.Path, default=ROOT / 'tools/entity_schema.py')
parser.add_argument('--require-dotnet', action='store_true')
OPTIONS, remaining = parser.parse_known_args()
spec = importlib.util.spec_from_file_location('content_tool', OPTIONS.tool)
tool = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tool)


class ContentClientTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.directory = pathlib.Path(temporary.name).resolve()
        self.schema = tool.load_schema(ROOT / 'schemas/player.schema.json')
        self.definition = self.directory / 'player.json'
        self.definition.write_text(json.dumps(self.schema), encoding='utf-8')

    def workbook(self, rows, second=False):
        book = openpyxl.Workbook()
        book.active.title = 'players'
        for row in rows: book.active.append(row)
        if second: book.create_sheet('ignored')
        path = self.directory / 'data.xlsx'
        book.save(path)
        book.close()
        return path

    def run_tool(self, *arguments):
        return subprocess.run([sys.executable, str(OPTIONS.tool), str(self.definition), *map(str, arguments)],
                              capture_output=True, text=True, encoding='utf-8')

    def test_csv_excel_equivalence_and_text_int64(self):
        path = self.workbook([['id', 'level', 'gold', 'speed', 'online'],
                              ['英雄', 3, str(2**63 - 1), 1.25, True], [None] * 5,
                              ['other', '4', 42, '2.5', 'false']])
        excel = tool.load_xlsx(path, self.schema)
        csv = self.directory / 'data.csv'
        csv.write_text('id,level,gold,speed,online\n英雄,3,9223372036854775807,1.25,true\nother,4,42,2.5,false\n', encoding='utf-8')
        self.assertEqual(excel, tool.load_csv(csv, self.schema))
        self.assertEqual(tool.generate(self.schema, excel), tool.generate(self.schema, tool.load_csv(csv, self.schema)))

    def test_excel_cells_are_strict_and_located(self):
        for column, value, diagnostic in [('gold', 10**15, 'text cell'), ('gold', 1.5, 'integer'),
                                          ('gold', True, 'integer'), ('gold', '9223372036854775808', 'expected integer'),
                                          ('level', 101, 'above maximum'), ('level', None, 'invalid integer'),
                                          ('online', 1, 'boolean'), ('online', 'TRUE', 'true or false'),
                                          ('id', 12, 'string'), ('secret', datetime.datetime(2026, 1, 1), 'date'),
                                          ('speed', '=1+1', 'formula'), ('secret', '#REF!', 'error')]:
            path = self.workbook([['id', column] if column != 'id' else ['secret', 'id'],
                                  ['hero', value] if column != 'id' else ['', value]])
            with self.subTest(column=column, value=value), self.assertRaisesRegex(ValueError, r'players!B2.*' + diagnostic):
                tool.load_xlsx(path, self.schema)
        for rows, diagnostic in [([['id', 'id'], ['a', 'a']], 'duplicate XLSX'),
                                 ([['id', 'typo'], ['a', 1]], 'unknown XLSX'),
                                 ([['gold'], [1]], 'missing id'), ([['id'], ['a'], ['a']], 'duplicate content ID'),
                                 ([['id'], [None, 'outside']], 'outside header')]:
            with self.subTest(rows=rows), self.assertRaisesRegex(ValueError, diagnostic):
                tool.load_xlsx(self.workbook(rows), self.schema)

    def test_sheet_selection_and_merges(self):
        path = self.workbook([['id'], ['hero']], second=True)
        with self.assertRaisesRegex(ValueError, 'multiple worksheets'): tool.load_xlsx(path, self.schema)
        self.assertEqual(tool.load_xlsx(path, self.schema, sheet='players')[0]['id'], 'hero')
        with self.assertRaisesRegex(ValueError, 'missing worksheet'): tool.load_xlsx(path, self.schema, sheet='absent')
        book = openpyxl.load_workbook(path)
        book['players'].merge_cells('A3:B3'); book.save(path); book.close()
        with self.assertRaisesRegex(ValueError, 'merged cells'): tool.load_xlsx(path, self.schema, sheet='players')

    def test_mixed_catalog_references_and_dependencies(self):
        schema = copy.deepcopy(self.schema)
        schema['fields'][3]['reference'] = {'table': 'items', 'allow_empty': False}
        self.definition.write_text(json.dumps(schema), encoding='utf-8')
        items = copy.deepcopy(self.schema); items['table'] = 'items'
        item_schema = self.directory / 'items.json'; item_schema.write_text(json.dumps(items))
        csv = self.directory / 'items.csv'; csv.write_text('id\nsword\n')
        xlsx = self.workbook([['id', 'secret'], ['hero', 'sword']], second=True)
        catalog = self.directory / 'catalog.json'
        catalog.write_text(json.dumps({'tables': [{'schema': 'player.json', 'xlsx': 'data.xlsx', 'sheet': 'players'},
                                                 {'schema': 'items.json', 'csv': 'items.csv'}]}))
        normalized = tool.load_schema(self.definition)
        self.assertEqual(tool.load_content(self.definition, normalized, catalog_path=catalog)[0]['secret'], 'sword')
        self.assertEqual(tool.catalog_inputs(catalog), [(self.definition, xlsx), (item_schema, csv)])
        self.workbook([['id', 'secret'], ['hero', 'missing']], second=True)
        with self.assertRaisesRegex(ValueError, r'players!B2.*unresolved reference'):
            tool.load_content(self.definition, normalized, catalog_path=catalog)

    def test_cli_preserves_both_outputs_and_checks_evolution(self):
        header, client = self.directory / 'generated.h', self.directory / 'client.cs'
        header.write_text('previous cpp'); client.write_text('previous cs')
        xlsx = self.workbook([['id', 'level'], ['hero', 999]])
        args = ['--xlsx', xlsx, '--output', header, '--csharp-output', client]
        self.assertNotEqual(self.run_tool(*args).returncode, 0)
        self.assertEqual(header.read_text(), 'previous cpp'); self.assertEqual(client.read_text(), 'previous cs')
        self.workbook([['id', 'level'], ['hero', 5]])
        previous = self.directory / 'previous.json'; previous.write_text(json.dumps(self.schema))
        self.schema['fields'][1]['default'] = 2
        self.definition.write_text(json.dumps(self.schema))
        result = self.run_tool(*args, '--previous', previous)
        self.assertIn('new version', result.stderr)
        self.assertEqual(client.read_text(), 'previous cs')
        self.schema['version'] += 1; self.definition.write_text(json.dumps(self.schema))
        result = self.run_tool(*args, '--previous', previous)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('ClientEntity', client.read_text())
        self.assertNotIn('hero', client.read_text())
        self.assertNotEqual(self.run_tool('--sheet', 'players', '--output', header).returncode, 0)
        self.assertNotEqual(self.run_tool('--output', header, '--csharp-output', header).returncode, 0)

    def test_compiled_csharp_contract(self):
        if not shutil.which('dotnet'):
            if OPTIONS.require_dotnet: self.fail('dotnet SDK is required')
            self.skipTest('dotnet SDK not installed')
        schema = copy.deepcopy(self.schema)
        schema['fields'][3]['default'] = 'SERVER_SECRET_SENTINEL'
        edge = tool.load_schema(ROOT / 'tests/schema/edge.schema.json')
        edge['fields'][1]['default'] += '\U0001f600\r\n"\\'
        # 'event' is valid in C++ and reserved in C#.
        edge['fields'][1]['name'] = 'event'
        edge['fields'].append(dict(id=9, name='maximum', type='int64', default=2**63-1, stored=True, visibility='owner'))
        for name, data in [('player', schema), ('edge', edge)]:
            generated = tool.generate_csharp(data)
            self.assertNotIn('SERVER_SECRET_SENTINEL', generated)
            self.assertNotIn('reference', generated)
            (self.directory / (name + '.cs')).write_text(generated, encoding='utf-8')
        (self.directory / 'ContractTest.csproj').write_text('''<Project Sdk="Microsoft.NET.Sdk">
  <PropertyGroup><OutputType>Exe</OutputType><TargetFramework>net8.0</TargetFramework>
  <TreatWarningsAsErrors>true</TreatWarningsAsErrors></PropertyGroup>
</Project>''')
        (self.directory / 'Program.cs').write_text('''using System;
using System.Text.Json;
using Player = Chwell.Generated.SchemaPlayer.ClientEntity;
using Contract = Chwell.Generated.SchemaPlayer.Contract;
class Program {
    static void Check(bool ok) { if (!ok) throw new Exception("Contract mismatch"); }
    static void Main() {
        var player = new Player();
        Check(player.value_id == "" && player.value_level == 1L && player.value_gold == 0L);
        Check(player.value_speed == 1D && !player.value_online);
        player.value_level = 6; Check(player.value_level == 6);
        Check(typeof(Player).GetProperties().Length == 5);
        Check(Contract.SchemaVersion == 1u && Contract.KeyFieldId == 1u && Contract.field_gold == 20u);
        using (var json = JsonDocument.Parse(Contract.SchemaJson)) {
            Check(json.RootElement.GetProperty("fields").GetArrayLength() == 5);
            Check(json.RootElement.GetProperty("reserved_ids")[0].GetUInt32() == 4u);
            foreach (var field in json.RootElement.GetProperty("fields").EnumerateArray()) {
                Check(field.GetProperty("visibility").GetString() != "server");
                Check(!field.TryGetProperty("reference", out var ignored));
            }
        }
        var edge = new Chwell.Generated.std.ClientEntity();
        Check(edge.value_count == long.MinValue && edge.value_maximum == long.MaxValue);
        Check(edge.value_event.Contains("\\0中文") && edge.value_event.EndsWith("😀\\r\\n\\\"\\\\"));
        Console.WriteLine("PASS: compiled C# client contract");
    }
}''', encoding='utf-8')
        result = subprocess.run(['dotnet', 'run', '--project', str(self.directory / 'ContractTest.csproj'),
                                 '--configuration', 'Release'], capture_output=True, text=True,
                                encoding='utf-8', errors='replace', timeout=120)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertIn('PASS:', result.stdout)


if __name__ == '__main__':
    unittest.main(argv=[sys.argv[0], *remaining])
